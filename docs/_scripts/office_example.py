"""Write ``docs/slides.pptx`` and ``docs/report.docx``, which the PowerPoint
and Word gallery reads.

python-pptx writes a chart the way PowerPoint saves one: a chart part with a
cached copy of each series' values, and the data in a workbook embedded beside
it. Its workbook leaves the cell above the categories empty, as PowerPoint's
own data sheet does, so each chart's workbook is written again with a header
there, as a reader who typed one into the data sheet would have it.
python-docx has no charts, so the document's are made the same way and moved
into it, with the relationships and content types Word writes for a chart.
Run from the repository root after changing it::

    uv run --with python-pptx --with python-docx python docs/_scripts/office_example.py
"""

from __future__ import annotations

import io
import posixpath
import re
import zipfile
from datetime import datetime
from pathlib import Path

import docx
import xlsxwriter
from docx.shared import Pt
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Cm

DOCS = Path(__file__).parents[1]
QUARTERS = ["Q1", "Q2", "Q3", "Q4"]
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CHART_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _chart(
    shapes: object,
    kind: object,
    box: tuple[int, int, int, int],
    header: str,
    categories: list[str],
    series: list[tuple[str, tuple[float, ...]]],
    number_format: str = "General",
) -> object:
    """Add a chart whose data sheet names its categories with ``header``."""
    data = CategoryChartData(number_format=number_format)
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    frame = shapes.add_chart(kind, *box, data)  # type: ignore[attr-defined]

    out = io.BytesIO()
    book = xlsxwriter.Workbook(out, {"in_memory": True})
    sheet = book.add_worksheet("Sheet1")
    numbers = book.add_format({"num_format": number_format})
    sheet.write(0, 0, header)
    for r, category in enumerate(categories, 1):
        sheet.write(r, 0, category)
    for c, (name, values) in enumerate(series, 1):
        sheet.write(0, c, name)
        for r, value in enumerate(values, 1):
            sheet.write_number(r, c, value, numbers)
    book.close()
    frame.chart.part.chart_workbook.update_from_xlsx_blob(out.getvalue())
    return frame


def _titled(chart: object, title: str, value_axis: str | None = None) -> None:
    chart.has_title = True  # type: ignore[attr-defined]
    chart.chart_title.text_frame.text = title  # type: ignore[attr-defined]
    if value_axis is not None:
        axis = chart.value_axis  # type: ignore[attr-defined]
        axis.has_title = True
        axis.axis_title.text_frame.text = value_axis


def _slide(deck: Presentation, title: str) -> object:
    slide = deck.slides.add_slide(deck.slide_layouts[5])  # Title Only
    slide.shapes.title.text = title
    return slide


def presentation() -> Presentation:
    """Quarterly sales over four slides, the last one hidden."""
    deck = Presentation()
    deck.core_properties.title = "Quarterly sales"
    deck.core_properties.created = datetime(2026, 1, 1)
    deck.core_properties.modified = datetime(2026, 1, 1)
    opening = deck.slides.add_slide(deck.slide_layouts[0])
    opening.shapes.title.text = "Quarterly sales"
    opening.placeholders[1].text = "North and South regions, 2025"

    box = (Cm(2), Cm(4), Cm(21), Cm(13))
    frame = _chart(
        _slide(deck, "Revenue").shapes,
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        box,
        "Quarter",
        QUARTERS,
        [("North", (120, 150, 90, 175)), ("South", (95, 110, 130, 160))],
    )
    frame.name = "Revenue chart"
    _titled(frame.chart, "Revenue by region", "Revenue ($k)")
    frame.chart.has_legend = True
    frame.chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    frame.chart.legend.include_in_layout = False

    frame = _chart(
        _slide(deck, "Margin").shapes,
        XL_CHART_TYPE.LINE_MARKERS,
        box,
        "Quarter",
        QUARTERS,
        [("Margin", (0.18, 0.21, 0.12, 0.24))],
        number_format="0%",
    )
    frame.name = "Margin chart"
    _titled(frame.chart, "Margin by quarter", "Margin")

    backup = _slide(deck, "Backup: revenue share")
    frame = _chart(
        backup.shapes,
        XL_CHART_TYPE.PIE,
        (Cm(4), Cm(4), Cm(17), Cm(13)),
        "Region",
        ["North", "South"],
        [("Share of revenue", (0.55, 0.45))],
        number_format="0%",
    )
    frame.name = "Share chart"
    _titled(frame.chart, "Share of revenue")
    # A slide kept for questions: hidden in the slide show.
    backup._element.set("show", "0")  # type: ignore[attr-defined]
    return deck


def _document_charts() -> Presentation:
    """The document's two charts, made by python-pptx to be moved."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])  # Blank

    box = (0, 0, Cm(15), Cm(8))
    chart = _chart(
        slide.shapes,
        XL_CHART_TYPE.LINE,
        box,
        "Month",
        ["Jan", "Feb", "Mar", "Apr", "May", "Jun"],
        [("Visitors", (1200, 1350, 1600, 1580, 1900, 2250))],
    ).chart
    _titled(chart, "Visitors by month", "Visitors")
    chart.has_legend = False

    chart = _chart(
        slide.shapes,
        XL_CHART_TYPE.BAR_CLUSTERED,
        box,
        "Team",
        ["Support", "Billing", "Sales"],
        [("Tickets closed", (340, 120, 85))],
    ).chart
    _titled(chart, "Tickets closed by team")
    chart.has_legend = False
    return deck


def _drawing(rid: str, number: int, name: str) -> str:
    """A chart drawn inline, as Word writes one."""
    width, height = Cm(15), Cm(8)
    return (
        "<w:r><w:drawing>"
        '<wp:inline distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="{width}" cy="{height}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:docPr id="{number}" name="{name}"/><wp:cNvGraphicFramePr/>'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        f'xmlns:r="{REL}" r:id="{rid}"/>'
        "</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>"
    )


def document() -> bytes:
    """A short report with two charts, each in a paragraph of its own."""
    report = docx.Document()
    report.core_properties.title = "Quarterly report"
    report.core_properties.created = datetime(2026, 1, 1)
    report.core_properties.modified = datetime(2026, 1, 1)
    report.styles["Normal"].font.size = Pt(11)
    report.add_heading("Quarterly report", level=1)
    report.add_paragraph(
        "Visitors grew every month but April, and passed two thousand in June."
    )
    report.add_paragraph("[chart 1]")
    report.add_paragraph("Support closed the most tickets of the three teams.")
    report.add_paragraph("[chart 2]")
    out = io.BytesIO()
    report.save(out)

    charts = io.BytesIO()
    _document_charts().save(charts)
    with zipfile.ZipFile(charts) as source:
        chart_parts = sorted(
            (
                n
                for n in source.namelist()
                if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)
            ),
            key=lambda n: int(re.search(r"\d+", posixpath.basename(n)).group()),
        )
        moved = {}
        for part in chart_parts:
            name = posixpath.basename(part)
            moved[f"word/charts/{name}"] = source.read(part)
            rels = f"ppt/charts/_rels/{name}.rels"
            moved[f"word/charts/_rels/{name}.rels"] = source.read(rels)
            for target in re.findall(
                rb'Target="\.\./embeddings/([^"]+)"',
                moved[f"word/charts/_rels/{name}.rels"],
            ):
                embedded = target.decode()
                moved[f"word/embeddings/{embedded}"] = source.read(
                    f"ppt/embeddings/{embedded}"
                )

    final = io.BytesIO()
    base = zipfile.ZipFile(out)
    target = zipfile.ZipFile(final, "w", zipfile.ZIP_DEFLATED)
    with base, target:
        for item in base.infolist():
            data = base.read(item.filename)
            if item.filename == "word/document.xml":
                text = data.decode("utf-8")
                for n, part in enumerate(chart_parts, 1):
                    paragraph = re.search(
                        rf"<w:r>(?:(?!<w:r>).)*?\[chart {n}\]</w:t></w:r>", text, re.S
                    )
                    assert paragraph is not None, f"no placeholder for chart {n}"
                    text = text.replace(
                        paragraph.group(0), _drawing(f"rIdChart{n}", n, f"Chart {n}")
                    )
                if "xmlns:wp=" not in text.split(">", 2)[1]:
                    raise SystemExit("python-docx no longer declares the wp namespace")
                data = text.encode("utf-8")
            elif item.filename == "word/_rels/document.xml.rels":
                extra = "".join(
                    f'<Relationship Id="rIdChart{n}" Type="{REL}/chart" '
                    f'Target="charts/{posixpath.basename(part)}"/>'
                    for n, part in enumerate(chart_parts, 1)
                )
                data = data.replace(
                    b"</Relationships>", extra.encode() + b"</Relationships>"
                )
            elif item.filename == "[Content_Types].xml":
                extra = (
                    f'<Default Extension="xlsx" ContentType="{XLSX_CONTENT_TYPE}"/>'
                    + "".join(
                        f'<Override PartName="/word/charts/{posixpath.basename(part)}" '
                        f'ContentType="{CHART_CONTENT_TYPE}"/>'
                        for part in chart_parts
                    )
                )
                data = data.replace(b"</Types>", extra.encode() + b"</Types>")
            target.writestr(item, data)
        for name, data in moved.items():
            target.writestr(name, data)
    return final.getvalue()


def main() -> None:
    presentation().save(str(DOCS / "slides.pptx"))
    (DOCS / "report.docx").write_bytes(document())


if __name__ == "__main__":
    main()
